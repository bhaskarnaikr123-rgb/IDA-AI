
from dataclasses import dataclass, asdict
from pathlib import Path
from typing import Optional
import re

from docx import Document as DocxDocument
from pypdf import PdfReader


SUPPORTED_EXTENSIONS = {".pdf", ".txt", ".docx"}


class DocumentIngestionError(Exception):
    """Raised when a document cannot be ingested safely."""


@dataclass
class DocumentSection:
    page_number: Optional[int]
    text: str


@dataclass
class IngestedDocument:
    filename: str
    file_type: str
    sections: list[DocumentSection]

    def to_dict(self) -> dict:
        return {
            "filename": self.filename,
            "file_type": self.file_type,
            "sections": [asdict(section) for section in self.sections],
        }


def clean_text(text: str) -> str:
    """Normalize whitespace without joining every word together."""
    text = text.replace("\x00", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def extract_pdf(path: Path) -> list[DocumentSection]:
    sections = []

    try:
        reader = PdfReader(str(path))

        if reader.is_encrypted:
            raise DocumentIngestionError(
                "Encrypted PDFs are not supported in this version."
            )

        for page_number, page in enumerate(reader.pages, start=1):
            text = clean_text(page.extract_text() or "")

            if text:
                sections.append(
                    DocumentSection(
                        page_number=page_number,
                        text=text,
                    )
                )

    except DocumentIngestionError:
        raise
    except Exception as exc:
        raise DocumentIngestionError(
            f"Could not read PDF: {exc}"
        ) from exc

    return sections


def extract_txt(path: Path) -> list[DocumentSection]:
    try:
        text = path.read_text(encoding="utf-8-sig")
    except UnicodeDecodeError as exc:
        raise DocumentIngestionError(
            "TXT file must use UTF-8 encoding."
        ) from exc
    except OSError as exc:
        raise DocumentIngestionError(
            f"Could not read TXT file: {exc}"
        ) from exc

    text = clean_text(text)

    return [
        DocumentSection(page_number=None, text=text)
    ] if text else []


def extract_docx(path: Path) -> list[DocumentSection]:
    try:
        document = DocxDocument(str(path))

        parts = [
            paragraph.text.strip()
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]

        # Include text stored in Word tables too.
        for table in document.tables:
            for row in table.rows:
                cells = [
                    cell.text.strip()
                    for cell in row.cells
                ]
                parts.append(" | ".join(cells))

        text = clean_text("\n".join(parts))

    except Exception as exc:
        raise DocumentIngestionError(
            f"Could not read DOCX file: {exc}"
        ) from exc

    return [
        DocumentSection(page_number=None, text=text)
    ] if text else []


def ingest_document(file_path: str | Path) -> IngestedDocument:
    """Extract text and metadata from a PDF, TXT, or DOCX file."""
    path = Path(file_path)

    if not path.exists() or not path.is_file():
        raise DocumentIngestionError(
            f"File does not exist: {path}"
        )

    extension = path.suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise DocumentIngestionError(
            f"Unsupported file type: {extension or '(no extension)'}. "
            "Supported types are PDF, TXT, and DOCX."
        )

    if extension == ".pdf":
        sections = extract_pdf(path)
    elif extension == ".txt":
        sections = extract_txt(path)
    else:
        sections = extract_docx(path)

    if not sections:
        raise DocumentIngestionError(
            "No extractable text found. The document may be empty, "
            "image-only, or unsupported."
        )

    return IngestedDocument(
        filename=path.name,
        file_type=extension.lstrip("."),
        sections=sections,
    )
