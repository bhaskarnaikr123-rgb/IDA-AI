
from pathlib import Path

import pytest
from docx import Document as DocxDocument

from app.ingestion import (
    DocumentIngestionError,
    ingest_document,
)


def test_ingest_txt(tmp_path: Path):
    file_path = tmp_path / "sample.txt"
    file_path.write_text(
        "Retrieval-Augmented Generation combines "
        "retrieval with language models.",
        encoding="utf-8",
    )

    result = ingest_document(file_path)

    assert result.filename == "sample.txt"
    assert result.file_type == "txt"
    assert len(result.sections) == 1
    assert result.sections[0].page_number is None
    assert "language models" in result.sections[0].text


def test_ingest_docx_including_tables(tmp_path: Path):
    file_path = tmp_path / "sample.docx"

    document = DocxDocument()
    document.add_paragraph("RAG retrieves relevant context.")

    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Component"
    table.cell(0, 1).text = "Purpose"

    document.save(file_path)

    result = ingest_document(file_path)
    text = result.sections[0].text

    assert result.file_type == "docx"
    assert "retrieves relevant context" in text
    assert "Component" in text
    assert "Purpose" in text


def test_reject_unsupported_file(tmp_path: Path):
    file_path = tmp_path / "sample.csv"
    file_path.write_text("name,value\nA,10", encoding="utf-8")

    with pytest.raises(DocumentIngestionError):
        ingest_document(file_path)


def test_reject_missing_file(tmp_path: Path):
    with pytest.raises(DocumentIngestionError):
        ingest_document(tmp_path / "missing.txt")


def test_reject_empty_txt(tmp_path: Path):
    file_path = tmp_path / "empty.txt"
    file_path.write_text("", encoding="utf-8")

    with pytest.raises(DocumentIngestionError):
        ingest_document(file_path)
